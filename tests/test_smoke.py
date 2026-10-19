#!/usr/bin/env python3
"""Smoke test for the MP2 pipeline.

    python tests/test_smoke.py outputs

Checks the cleaning counts, sample sizes, the fit of the preferred model (M6) and that every visual output exists.
"""
import json
import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "outputs"
with open(os.path.join(OUT, "results.json"), encoding="utf-8") as fh:
    R = json.load(fh)

c = R["clean"]
assert c["n_raw"] == 526, c["n_raw"]
assert c["n_flag_cells"] == 49 and c["n_rows_flagged"] == 43 and c["n_recovered"] == 11
assert c["n_final"] == 502 and c["n_final_fam"] == 491 and c["n_ten_gt_exper"] == 13
assert abs(R["fit"]["M6"]["r2"] - 0.525283) < 1e-5, R["fit"]["M6"]["r2"]
assert abs(R["fit"]["M6"]["adj_r2"] - 0.501398) < 1e-5
assert abs(R["fit"]["M6"]["aic"] - 544.848) < 1e-2
assert 0 < R["diag"]["gq_exper"]["p"] < 1

figs = [f for f in os.listdir(os.path.join(OUT, "figures")) if f.endswith(".png")]
tabs = [f for f in os.listdir(os.path.join(OUT, "table_images")) if f.endswith(".png")]
assert len(figs) == 18, f"expected 18 figures, found {len(figs)}"
assert "00_summary_dashboard.png" in figs
assert len(tabs) == 14, f"expected 14 table images, found {len(tabs)}"

from openpyxl import load_workbook
wb = load_workbook(os.path.join(OUT, "MP2_results.xlsx"))
assert len(wb.sheetnames) == 19 and "regression_M1_M6" in wb.sheetnames, wb.sheetnames
print(f"smoke test passed: {len(figs)} figures, {len(tabs)} table images, {len(wb.sheetnames)} workbook sheets; "
      f"N = {c['n_final']}; M6 R2 = {R['fit']['M6']['r2']:.6f}")
