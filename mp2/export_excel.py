"""Write one formatted Excel workbook: overview + dashboard picture, every results table, numeric coefficients, data.

The tables are exactly the ones rendered as images (same Table objects), so the two can never disagree.
The workbook holds static results (no formulas): it is a presentation and look-up copy of the analysis.
"""
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .table_builders import build_all

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="2A78D6")
ZEBRA = PatternFill("solid", fgColor="F4F3EF")
THIN = Side(style="thin", color="8A8984")


def _title(ws, title, note=None):
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws.row_dimensions[1].height = 24
    if note:
        ws["A2"] = note
        ws["A2"].font = Font(name=FONT, size=9, italic=True, color="52514E")


def _header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEAD_FILL
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 34


def _widths(ws, df, header_row, cap=60, minimum=10):
    for j, col in enumerate(df.columns, start=1):
        longest = max([max(len(x) for x in str(col).split("\n"))] + [len(str(v)) for v in df.iloc[:, j - 1]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, minimum), cap)


def _table_sheet(wb, tab):
    ws = wb.create_sheet(tab.key[:31])
    _title(ws, tab.title, tab.note)
    hr = 4
    df = tab.df
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=hr, column=j, value=str(col))
    _header(ws, hr, len(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=hr + i, column=j, value=v)
            cell.font = Font(name=FONT, size=10, bold=(i - 1) in tab.bold_rows,
                             color="52514E" if (tab.group == 2 and str(row[0]) == "") else "000000")
            cell.alignment = Alignment(horizontal="left" if (j - 1) in tab.left_cols else "center", vertical="center")
            if ((i - 1) // tab.group) % 2 == 1:
                cell.fill = ZEBRA
    last = hr + len(df)
    for c in range(1, len(df.columns) + 1):
        ws.cell(row=last, column=c).border = Border(bottom=THIN)
    _widths(ws, df, hr)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    ws.sheet_view.showGridLines = False
    return ws


def _data_sheet(wb, name, df, title, note=None, decimals=4):
    ws = wb.create_sheet(name[:31])
    _title(ws, title, note)
    hr = 4
    for j, col in enumerate(df.columns, start=1):
        ws.cell(row=hr, column=j, value=str(col))
    _header(ws, hr, len(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start=hr + 1):
        for j, v in enumerate(row, start=1):
            if isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = None if np.isnan(v) else float(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            elif v is not None and isinstance(v, float) and np.isnan(v):
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name=FONT, size=10)
            if isinstance(v, float):
                cell.number_format = "0." + "0" * decimals
    for j, col in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(j)].width = min(max(len(str(col)) + 3, 11), 32)
    ws.freeze_panes = ws.cell(row=hr + 1, column=2)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(df.columns))}{hr + len(df)}"
    return ws


def _overview(wb, ctx, tables):
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    _title(ws, "Salary determinants at HCL Technologies: results workbook",
           f"OLS on ln(hourly salary), HC3-robust inference. N = {ctx.N}. Static results; the analysis code regenerates every number.")
    ws["A4"], ws["B4"] = "Sheet", "Contents"
    for c in (1, 2):
        ws.cell(row=4, column=c).fill = HEAD_FILL
        ws.cell(row=4, column=c).font = Font(name=FONT, bold=True, color="FFFFFF")
    items = [(t.key[:31], t.title) for t in tables] + [
        ("coefficients_numeric", "All model coefficients as numbers (b, s.e., t, p, 95% CI) for M1-M6 and the family model"),
        ("cleaning_log", "Every cleaning action: record id, variable, raw value, issue, action"),
        ("cleaned_data", "Full data after cleaning (missing values left blank)"),
        ("analysis_sample", "The N = %d employees used in the main models, with derived variables" % ctx.N)]
    for i, (s, d) in enumerate(items, start=5):
        c = ws.cell(row=i, column=1, value=s)
        c.hyperlink = f"#'{s}'!A1"
        c.font = Font(name=FONT, color="2A78D6", underline="single")
        ws.cell(row=i, column=2, value=d).font = Font(name=FONT)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    img = os.path.join(ctx.fig_dir, "00_summary_dashboard.png")
    if os.path.exists(img):
        pic = XLImage(img)
        scale = 1100 / pic.width
        pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
        ws.add_image(pic, f"A{len(items) + 7}")


def run(ctx):
    tables = getattr(ctx, "tables", None) or build_all(ctx)
    wb = Workbook()
    _overview(wb, ctx, tables)
    for tab in tables:
        _table_sheet(wb, tab)
    coefs = []
    for k, tb in ctx.TBL.items():
        t = tb[["term", "label", "b", "se_hc3", "t_hc3", "p_hc3", "lo", "hi"]].copy()
        t.insert(0, "model", k)
        coefs.append(t)
    fam = ctx.TB5F[["term", "label", "b", "se_hc3", "t_hc3", "p_hc3", "lo", "hi"]].copy()
    fam.insert(0, "model", "M4 + family")
    coefs.append(fam)
    _data_sheet(wb, "coefficients_numeric", pd.concat(coefs, ignore_index=True), "Model coefficients (numeric)",
                "HC3-robust s.e., t-distribution p-values and 95% confidence intervals.", decimals=5)
    _data_sheet(wb, "cleaning_log", ctx.log.drop(columns="cat"), "Cleaning log", "One row per cell changed or recovered.")
    cleaned = ctx.df.drop(columns=["complete_core", "complete_fam"])
    _data_sheet(wb, "cleaned_data", cleaned, "Cleaned data", "Impossible values removed; see cleaning_log.", decimals=4)
    _data_sheet(wb, "analysis_sample", ctx.S.drop(columns=[c for c in ("y_bc",) if c in ctx.S.columns]), "Analysis sample (main models)",
                "Includes residuals, leverage, Cook's distance and studentised residuals from M6.", decimals=4)
    wb.save(os.path.join(ctx.out, "MP2_results.xlsx"))
